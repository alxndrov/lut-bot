import asyncio
import io
import json
import tempfile
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

import aiosqlite
from openpyxl import load_workbook
import config
import database as db
from services import tax_account as tax, finance_reports as reports, payout
from handlers import finance, reviews


class ScheduleTests(unittest.TestCase):
    def test_leap_and_short_months_and_year_change(self):
        for year, month, day, want in [(2024, 2, 15, 29), (2025, 2, 15, 28),
                                      (2026, 4, 15, 30), (2026, 1, 15, 31)]:
            now = datetime(year, month, day, config.DAILY_REPORT_HOUR_MSK, tzinfo=reports.MSK)
            self.assertEqual(reports.next_boundary(now).day, want)
        now = datetime(2026, 12, 31, config.DAILY_REPORT_HOUR_MSK, tzinfo=reports.MSK)
        self.assertEqual(reports.next_boundary(now).strftime('%Y-%m-%d'), '2027-01-15')
        self.assertEqual(reports.previous_boundary(reports.next_boundary(now)), now)

    def test_tax_parser_including_zero(self):
        for text, value in [('0', 0), ('0,00', 0), ('1 250,50', 1250.5), ('12.01', 12.01)]:
            self.assertEqual(tax.parse_amount(text), value)
        for text in ['-1', 'nan', 'inf', '1.001', '12 abc', '1 2', '']:
            with self.assertRaises(ValueError):
                tax.parse_amount(text)

    def test_menus(self):
        callbacks = [b.callback_data for row in finance._finance_keyboard().inline_keyboard for b in row]
        self.assertNotIn('rev_show', callbacks)
        self.assertTrue({'fin_cash', 'fin_revenue', 'fin_reports', 'cash_log'}.issubset(callbacks))
        self.assertNotIn('ops_menu', callbacks)
        self.assertNotIn('ops_menu', [b.callback_data for row in reviews._keyboard(0, 0, []).inline_keyboard for b in row])


class DatabaseTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db_patch = patch.object(db, 'DB_PATH', str(Path(self.tmp.name) / 'test.db'))
        self.db_patch.start()
        await db.init_db()
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute("INSERT INTO products(id, name, description, price, category) VALUES (1, 'Test', '', 1000, 'physical')")
            await conn.commit()

    async def asyncTearDown(self):
        self.db_patch.stop()
        self.tmp.cleanup()

    async def purchase(self, amount, when):
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute('INSERT INTO purchases(user_id, product_id, amount, created_at) VALUES (1, 1, ?, ?)', (amount, when))
            await conn.commit()

    async def test_manual_tax_partial_payment_and_concurrent_duplicate_clicks(self):
        await self.purchase(10000, '2026-09-05 10:00:00')
        await tax.edit('2026-09', 250, 1)
        self.assertEqual((await tax.assessment('2026-09'))['accrued'], 250)
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute("INSERT INTO npd_payments(amount, tax_month) VALUES (100, '2026-09')")
            await conn.commit()
        await tax.edit('2026-09', 300, 1)
        payments = await asyncio.gather(tax.mark_paid('2026-09', 1, 'Test'), tax.mark_paid('2026-09', 1, 'Test'))
        self.assertEqual(sum(p[1] for p in payments), 200)
        self.assertEqual((await tax.assessment('2026-09'))['paid'], 300)
        with self.assertRaises(ValueError):
            await tax.edit('2026-09', 500, 1)
        # Payment date is the actual date; allocation stays with the tax month.
        paid = await db.get_npd_payments_by_month(100)
        self.assertEqual(paid[0]['month'], '2026-09')
        self.assertEqual(paid[0]['total'], 300)

    async def test_tax_allocation_split_and_zero(self):
        await self.purchase(1000, '2026-09-05 10:00:00')
        await self.purchase(3000, '2026-09-20 10:00:00')
        await tax.edit('2026-09', 80, 1)
        first = await tax.accrued_for_period('2026-09-01', '2026-09-14')
        second = await tax.accrued_for_period('2026-09-15', '2026-09-30')
        self.assertEqual((first, second), (20, 60))
        split = await payout.split('2026-09-01', '2026-09-30')
        self.assertEqual(split['npd'], 80)
        self.assertAlmostEqual(split['owner'] + split['partner'], split['net'])
        self.assertAlmostEqual(split['partner'], split['net'] * config.PARTNER_GOODS_PERCENT_NEW / 100)
        await tax.edit('2026-09', 0, 1)
        self.assertEqual(await tax.mark_paid('2026-09', 1, 'Test'), (None, 0))
        await tax.edit('2026-09', 40, 1)  # zero is still editable
        with self.assertRaises(ValueError):
            await tax.edit('2026-09', float('nan'), 1)

    async def test_old_shares_and_payment_summary_use_corrected_tax(self):
        from services.daily_report import fetch_payments_summary
        await self.purchase(1000, '2026-08-05 10:00:00')
        await tax.edit('2026-08', 20, 1)
        split = await payout.split('2026-08-01', '2026-08-31')
        expected_net = 1000 * (1 - config.PRODAMUS_FEE_PERCENT / 100) - 20
        self.assertAlmostEqual(split['net'], expected_net)
        self.assertAlmostEqual(split['partner'], expected_net * config.PARTNER_GOODS_PERCENT / 100)
        summary = await fetch_payments_summary('2026-08-01T00:00:00.000Z', '2026-08-31T23:59:59.999Z')
        self.assertEqual(summary['npd'], 20)
        self.assertAlmostEqual(summary['net'], expected_net)

    async def test_report_without_sales_keeps_expenses_and_actual_cdek(self):
        end = datetime(2026, 9, 30, config.DAILY_REPORT_HOUR_MSK, tzinfo=reports.MSK)
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute("INSERT INTO expenses(category, amount, spent_at) VALUES ('Пластик', 200, '2026-09-20 10:00:00')")
            await conn.execute("INSERT INTO cdek_payments(amount, paid_at) VALUES (300, '2026-09-20 10:00:00')")
            await conn.commit()
        data = await reports.build(end)
        self.assertEqual(data['gross'], 0)
        self.assertEqual(data['expenses'], 200)
        self.assertEqual(data['cdek'], 300)
        self.assertEqual(data['net'], -200)  # payment settles prior delivery, not a new expense
        self.assertAlmostEqual(data['partner'] + data['owner'], -200)

    async def test_tax_month_msk_and_migration_idempotence(self):
        await self.purchase(1000, '2026-08-31 21:00:00')
        await db.add_npd_payment(20, paid_at='2026-08-28 12:00:00')
        await db.init_db()
        await db.init_db()
        monthly = {r['month']: r for r in await tax.months()}
        self.assertEqual(monthly['2026-09']['gross'], 1000)
        self.assertEqual(monthly['2026-08']['paid'], 20)

    async def test_tax_payment_summary_uses_actual_date(self):
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute("""INSERT INTO npd_payments
                (amount, paid_at, tax_month, comment) VALUES (?, ?, ?, ?)""",
                (4867, '2026-09-14 06:22:56', '2026-08', 'НПД за 2026-08'))
            await conn.commit()

        september = await db.get_npd_payments_summary('2026-09-01 00:00:00', '2026-09-30 23:59:59')
        august = await db.get_npd_payments_summary('2026-08-01 00:00:00', '2026-08-31 23:59:59')
        self.assertEqual(september['total'], 4867)
        self.assertEqual(august['total'], 0)

    async def test_cash_keeps_current_tax_reserve_when_old_tax_paid_today(self):
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute("""INSERT INTO npd_payments
                (amount, paid_at, tax_month, comment) VALUES (?, ?, ?, ?)""",
                (4867, '2026-09-14 06:22:56', '2026-08', 'НПД за 2026-08'))
            await conn.commit()

        s = {
            'gross': 22591, 'fee': 858.458, 'expenses': 3803,
            'delivery_out': 5165.33, 'npd': 903.64, 'fee_pct': config.PRODAMUS_FEE_PERCENT,
        }
        # Резервы и «свободно» — по всему счёту (как в листе «Финансы»),
        # а не по периоду
        account = {'on_account': 17929.54, 'pending': 0, 'cdek_reserve': 5165.33,
                   'tax_reserve': 903.64, 'free': 11860.57}
        with patch.object(finance, '_free_cash', AsyncMock(return_value=account)):
            text = await finance._cash_block_text(
                s, '2026-09-02 08:47:58', '2026-09-14 06:39:21',
                {'total': 13123.47, 'count': 3},
                await db.get_npd_payments_summary('2026-09-02 08:47:58', '2026-09-14 06:39:21'),
                0,
            )
        self.assertIn('Оплачено НПД по факту: −4,867.00 ₽', text)
        self.assertIn('СДЭК 5,165.33 ₽ · НПД 903.64 ₽', text)
        self.assertIn('Свободно к выплате: 11,860.57 ₽', text)

    def test_payable_never_exceeds_free_cash(self):
        s = {'owner_left': 600, 'partner_left': 400}
        self.assertEqual(finance._payable(s, 5000), (600, 400))
        self.assertEqual(finance._payable(s, 500), (300, 200))
        self.assertEqual(finance._payable(s, -323.2), (0, 0))
        self.assertEqual(finance._payable({'owner_left': -50, 'partner_left': 400}, 1000), (0, 400))

    async def test_tax_ui_old_months_and_no_payouts(self):
        for month in range(1, 9):
            await self.purchase(100, f'2025-{month:02d}-01 10:00:00')
        with patch.object(db, 'get_payouts', side_effect=AssertionError('must not load payouts')):
            text, keyboard = await finance._cash_log_render()
        callbacks = [b.callback_data for row in keyboard.inline_keyboard for b in row]
        self.assertIn('npd_page:1', callbacks)
        self.assertNotIn('Выплаты', text)
        _, older = await finance._cash_log_render(1)
        self.assertIn('npd_detail:2025-01', [b.callback_data for row in older.inline_keyboard for b in row])

    async def test_tax_detail_does_not_list_purchases(self):
        purchases = [
            {
                'amount': 100,
                'created_at': '2026-09-01 10:00:00',
                'username': 'buyer_with_a_long_name',
                'first_name': None,
                'user_id': i,
                'product_name': 'Очень длинное название товара ' * 4,
            }
            for i in range(150)
        ]
        text = "\n".join(finance._npd_detail_lines('2026-09', purchases, None))
        self.assertLess(len(text), 4096)
        self.assertIn('Продаж: <b>150</b>', text)
        self.assertNotIn('Очень длинное название товара', text)
        self.assertNotIn('buyer_with_a_long_name', text)

    async def test_finance_open_does_not_fetch_data(self):
        message = SimpleNamespace(from_user=SimpleNamespace(id=123), answer=AsyncMock())
        with patch.object(config, 'ADMIN_IDS', [123]), patch.object(finance, '_settle_split', side_effect=AssertionError()), patch.object(finance, '_finance_pulse_text', side_effect=AssertionError()):
            await finance.cmd_finance(message, SimpleNamespace(clear=AsyncMock()))
        self.assertEqual(message.answer.await_count, 1)
        self.assertNotIn('₽', message.answer.call_args.args[0])

    async def test_report_boundary_and_immutable_xlsx(self):
        end = datetime(2026, 9, 15, config.DAILY_REPORT_HOUR_MSK, tzinfo=reports.MSK)
        start = reports.previous_boundary(end)
        await self.purchase(1000, reports.stamp(start))
        await self.purchase(2000, reports.stamp(end - timedelta(seconds=1)))
        await self.purchase(4000, reports.stamp(end))
        rid, data = await reports.save(end)
        self.assertEqual(data['gross'], 3000)
        self.assertEqual(data['count'], 2)
        second = await reports.build(reports.next_boundary(end))
        self.assertEqual(second['gross'], 4000)
        await self.purchase(9000, reports.stamp(start + timedelta(seconds=1)))
        self.assertEqual((await reports.save(end))[1], data)
        workbook = load_workbook(io.BytesIO(reports.xlsx(await reports.get(rid))), data_only=True)
        sheet = workbook.active
        self.assertEqual(sheet['B4'].value, 3000)
        for i, (label, value) in enumerate(reports.rows(data), start=4):
            self.assertEqual(sheet.cell(i, 1).value, label)
            self.assertEqual(sheet.cell(i, 2).value, round(value, 2))
        self.assertEqual(sheet['A10'].value, 'Доля Дани')
        self.assertEqual(sheet['A11'].value, 'Доля Миши')
        self.assertEqual(sheet['A15'].value, reports.note(data))

    async def test_scheduler_restart_catchup_and_recipient_retry(self):
        before = datetime(2026, 9, 13, 9, tzinfo=reports.MSK)
        bot = SimpleNamespace(send_message=AsyncMock())
        with patch.object(config, 'ADMIN_IDS', [1, 2]):
            await reports.tick(bot, before)
            bot.send_message.assert_not_awaited()
            async def send(admin, *args, **kwargs):
                if admin == 2:
                    raise RuntimeError('simulated Telegram failure')
            bot.send_message.side_effect = send
            with self.assertLogs(reports.logger, level='ERROR'):
                await reports.tick(bot, before + timedelta(days=2))
            self.assertEqual(bot.send_message.await_count, 2)
            bot.send_message.reset_mock()
            bot.send_message.side_effect = None
            await reports.tick(bot, before + timedelta(days=2))
            self.assertEqual(bot.send_message.await_count, 1)
            self.assertEqual(bot.send_message.call_args.args[0], 2)
            await reports.tick(bot, before + timedelta(days=2))
            self.assertEqual(bot.send_message.await_count, 1)
            await reports.tick(bot, datetime(2026, 10, 15, 9, tzinfo=reports.MSK))
            self.assertEqual(len(await reports.history()), 3)

    async def test_old_saved_schedule_moves_to_new_dates_without_early_send(self):
        bot = SimpleNamespace(send_message=AsyncMock())
        for old_day, new_day in [(14, 15), (29, 30)]:
            old = datetime(2026, 9, old_day, config.DAILY_REPORT_HOUR_MSK, tzinfo=reports.MSK)
            async with aiosqlite.connect(db.DB_PATH) as conn:
                await conn.execute('INSERT OR REPLACE INTO finance_report_schedule VALUES (1, ?)', (old.isoformat(),))
                await conn.commit()
            await reports.tick(bot, old)
            bot.send_message.assert_not_awaited()
            async with aiosqlite.connect(db.DB_PATH) as conn:
                async with conn.execute('SELECT next_at FROM finance_report_schedule WHERE id=1') as cur:
                    target = datetime.fromisoformat((await cur.fetchone())[0])
            self.assertEqual(target.day, new_day)
            self.assertEqual(target.hour, config.DAILY_REPORT_HOUR_MSK)
        self.assertEqual(await reports.history(), [])

    async def test_restricted_tax_and_export_handlers(self):
        cb = SimpleNamespace(from_user=SimpleNamespace(id=999), answer=AsyncMock(), message=SimpleNamespace(answer_document=AsyncMock()))
        with patch.object(config, 'ADMIN_IDS', [1]), patch.object(reports, 'get', side_effect=AssertionError()):
            await finance.cb_fin_xlsx(cb)
        cb.message.answer_document.assert_not_awaited()
        self.assertTrue(cb.answer.call_args.kwargs['show_alert'])


if __name__ == '__main__':
    unittest.main()

"""Saved financial statements at 09:00 MSK on the 15th and last day."""
import asyncio
import calendar
import io
import json
import logging
from datetime import datetime, timedelta, timezone
from html import escape

import aiosqlite
from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import config
import database as db
from services import payout

MSK = timezone(timedelta(hours=3))
logger = logging.getLogger(__name__)


def boundaries(now: datetime):
    now = now.astimezone(MSK)
    result = []
    for offset in (-1, 0, 1):
        index = now.year * 12 + now.month - 1 + offset
        year, m0 = divmod(index, 12)
        month = m0 + 1
        for day in (15, calendar.monthrange(year, month)[1]):
            result.append(datetime(year, month, day, config.DAILY_REPORT_HOUR_MSK, tzinfo=MSK))
    return sorted(result)


def next_boundary(now: datetime) -> datetime:
    return next(d for d in boundaries(now) if d > now)


def previous_boundary(now: datetime) -> datetime:
    return max(d for d in boundaries(now) if d < now)


def stamp(moment: datetime) -> str:
    return moment.astimezone(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')


async def build(end: datetime) -> dict:
    from handlers.finance import _pending_gross
    start = previous_boundary(end)
    # All DB helpers use inclusive bounds; schedule periods are [start, end).
    dt_from, dt_to = stamp(start), stamp(end - timedelta(seconds=1))
    share = await payout.split(dt_from, dt_to)
    cdek = await db.get_cdek_payments_summary(dt_from, dt_to)
    # Outstanding remittances are a balance as of the end, not a period expense.
    pending = await _pending_gross('2020-01-01 00:00:00', dt_to, end.astimezone(timezone.utc))
    return {
        'start': start.isoformat(), 'end': end.isoformat(),
        'gross': share['gross'], 'count': share['count'],
        'cdek': float(cdek['total']), 'tax': share['npd'],
        'expenses': share['expenses'], 'refunds': share.get('refunds', 0),
        'pending': pending * (1 - share['fee_pct'] / 100),
        'net': share['net'], 'partner': share['partner'], 'owner': share['owner'],
        'partner_name': config.PARTNER_NAME, 'owner_name': config.OWNER_NAME,
        'fee': share['fee'], 'delivery_accrued': share['delivery_out'],
    }


def _genitive(name):
    return name[:-1] + "и" if name.endswith(("а", "я")) else name


def rows(data):
    return [
        ('Принято платежей', data['gross']),
        *([('Возвращено покупателям', data['refunds'])] if data.get('refunds') else []),
        ('Потрачено на СДЭК', data['cdek']),
        ('Отложено на налог', data['tax']),
        ('Потрачено на расходники', data['expenses']),
        ('Ещё не перевёл Prodamus (оценка)', data['pending']),
        ('Чистая прибыль', data['net']),
        (f"Доля {_genitive(data['partner_name'])}", data['partner']),
        (f"Доля {_genitive(data['owner_name'])}", data['owner']),
    ]


def period_label(data):
    return ' — '.join(datetime.fromisoformat(data[k]).strftime('%d.%m.%Y %H:%M') for k in ('start', 'end')) + ' МСК'


def note(data):
    return (f"Прибыль = платежи − "
            + ("возвраты − " if data.get('refunds') else "")
            + f"комиссия Prodamus ({data['fee']:,.2f} ₽) "
            f"− налог − начисленная доставка ({data['delivery_accrued']:,.2f} ₽) − расходы. "
            'Оплаченный СДЭК показан по факту; начисленная доставка вычтена из прибыли, '
            'даже если счёт ещё не оплачен. Расходники — все внесённые расходы. '
            'Налог учитывает ручные исправления, распределённые по выручке месяца. '
            'Ожидаемый перевод Prodamus указан после комиссии, по срокам переводов '
            'без учёта праздников и сверки с банком. '
            'Конец периода не включён; этот момент открывает следующую сводку.')


def render(data):
    lines = ['📊 <b>Финансовая сводка</b>', period_label(data), '']
    lines.extend(f'{escape(label)}: <b>{value:,.2f} ₽</b>' for label, value in rows(data))
    lines += ['', f"Платежей: {data['count']}", f'<i>{note(data)}</i>']
    return '\n'.join(lines)


def keyboard(report_id):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text='📥 Скачать XLSX', callback_data=f'fin_xlsx:{report_id}')],
        [InlineKeyboardButton(text='📊 Все сводки', callback_data='fin_reports')],
        [InlineKeyboardButton(text='◀️ Назад', callback_data='fin_show')],
    ])


def xlsx(data):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Сводка'
    sheet.append(['Финансовая сводка'])
    sheet.append([period_label(data)])
    sheet.append(['Показатель', 'Сумма, ₽'])
    for label, value in rows(data):
        sheet.append([label, round(value, 2)])
        sheet.cell(sheet.max_row, 2).number_format = '#,##0.00'
    sheet.append(['Количество платежей', data['count']])
    sheet.append(['Комиссия Prodamus', round(data['fee'], 2)])
    sheet.append(['Начисленная доставка СДЭК', round(data['delivery_accrued'], 2)])
    for n in (13, 14):
        sheet.cell(n, 2).number_format = '#,##0.00'
    sheet.append([note(data)])
    sheet.merge_cells(start_row=15, start_column=1, end_row=15, end_column=2)
    sheet['A15'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[15].height = 130
    sheet.column_dimensions['A'].width = 65
    sheet.column_dimensions['B'].width = 24
    sheet.freeze_panes = 'B4'
    for cell in sheet[3]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='234B65')
    sheet['A1'].font = Font(bold=True, size=16)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


async def save(end: datetime):
    key = stamp(end)
    async with aiosqlite.connect(db.DB_PATH) as conn:
        async with conn.execute('SELECT id, data_json FROM finance_reports WHERE scheduled_at=?', (key,)) as cur:
            existing = await cur.fetchone()
    if existing:
        return existing[0], json.loads(existing[1])
    data = await build(end)
    async with aiosqlite.connect(db.DB_PATH) as conn:
        await conn.execute('INSERT OR IGNORE INTO finance_reports(scheduled_at, data_json) VALUES (?, ?)',
                           (key, json.dumps(data, ensure_ascii=False)))
        await conn.commit()
        async with conn.execute('SELECT id, data_json FROM finance_reports WHERE scheduled_at=?', (key,)) as cur:
            row = await cur.fetchone()
            return row[0], json.loads(row[1])


async def get(report_id):
    async with aiosqlite.connect(db.DB_PATH) as conn:
        async with conn.execute('SELECT data_json FROM finance_reports WHERE id=?', (report_id,)) as cur:
            row = await cur.fetchone()
            return json.loads(row[0]) if row else None


async def history(page=0):
    async with aiosqlite.connect(db.DB_PATH) as conn:
        async with conn.execute('SELECT id, data_json FROM finance_reports ORDER BY scheduled_at DESC LIMIT 11 OFFSET ?',
                                (max(0, page) * 10,)) as cur:
            return [(r[0], json.loads(r[1])) for r in await cur.fetchall()]


async def tick(bot, now: datetime):
    """Persist the next due time, catch up after downtime, retry failed recipients."""
    async with aiosqlite.connect(db.DB_PATH) as conn:
        await conn.execute('INSERT OR IGNORE INTO finance_report_schedule(id, next_at) VALUES (1, ?)',
                           (next_boundary(now).isoformat(),))
        await conn.commit()
        async with conn.execute('SELECT next_at FROM finance_report_schedule WHERE id=1') as cur:
            target = datetime.fromisoformat((await cur.fetchone())[0])
    # Move a persisted date from the former 14th/penultimate-day schedule
    # to the next new boundary, including when catching up after downtime.
    if target not in boundaries(target):
        target = next_boundary(target)
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute('UPDATE finance_report_schedule SET next_at=? WHERE id=1', (target.isoformat(),))
            await conn.commit()
    while target <= now:
        await save(target)
        target = next_boundary(target)
        async with aiosqlite.connect(db.DB_PATH) as conn:
            await conn.execute('UPDATE finance_report_schedule SET next_at=? WHERE id=1', (target.isoformat(),))
            await conn.commit()
    for admin_id in config.ADMIN_IDS:
        async with aiosqlite.connect(db.DB_PATH) as conn:
            async with conn.execute('''SELECT r.id, r.data_json FROM finance_reports r
                LEFT JOIN finance_report_deliveries d ON d.report_id=r.id AND d.admin_id=?
                WHERE d.report_id IS NULL ORDER BY r.scheduled_at''', (admin_id,)) as cur:
                pending = await cur.fetchall()
        for report_id, payload in pending:
            try:
                await bot.send_message(admin_id, render(json.loads(payload)), parse_mode='HTML',
                                       reply_markup=keyboard(report_id))
            except Exception:
                logger.exception('finance_reports: delivery failed for %s / %s', report_id, admin_id)
                continue
            async with aiosqlite.connect(db.DB_PATH) as conn:
                await conn.execute('INSERT OR IGNORE INTO finance_report_deliveries VALUES (?, ?)', (report_id, admin_id))
                await conn.commit()


async def report_loop(bot):
    while True:
        try:
            await tick(bot, datetime.now(MSK))
        except Exception:
            logger.exception('finance_reports: tick failed')
        await asyncio.sleep(30)
